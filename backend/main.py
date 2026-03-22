"""
NeuraNote Quiz Module - Backend
"""

from fastapi import FastAPI, File, UploadFile, Form, Depends, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import Text
from typing import List, Optional
from datetime import datetime
import json
import os
import io
from dotenv import load_dotenv

load_dotenv()

from config.database import get_db, engine
from models import quizmodels as models
from db import schemas

# AI/LangChain imports
from PyPDF2 import PdfReader
from langchain_core.prompts import PromptTemplate
from langchain_openai import ChatOpenAI
from langchain_text_splitters import RecursiveCharacterTextSplitter

# Document processing
import openpyxl
from docx import Document

# PDF generation
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER

app = FastAPI(title="NeuraNote Quiz API - With Levels")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

models.Base.metadata.create_all(bind=engine)

# Initialize LLM
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
llm = None
if OPENAI_API_KEY and OPENAI_API_KEY.startswith("sk-"):
    try:
        llm = ChatOpenAI(model="gpt-3.5-turbo", temperature=0.7, api_key=OPENAI_API_KEY)
        print("✅ LangChain LLM initialized successfully")
    except Exception as e:
        print(f"⚠️ LLM initialization failed: {e}")
else:
    print("⚠️ No valid OpenAI API key found - will use mock questions")


# ==================== HELPER FUNCTIONS ====================

# Difficulty progression mapping
DIFFICULTY_PROGRESSION = {
    'easy': 'medium',
    'medium': 'hard',
    'hard': None  # No next level after hard
}

def get_next_difficulty(current_difficulty: str) -> Optional[str]:
    """Get the next difficulty level"""
    return DIFFICULTY_PROGRESSION.get(current_difficulty.lower())


def extract_text_from_file(uploaded_file: UploadFile) -> str:
    """Extract text from uploaded files"""
    try:
        filename = uploaded_file.filename.lower()
        
        if filename.endswith(".pdf"):
            reader = PdfReader(uploaded_file.file)
            text = ""
            for page in reader.pages:
                text += page.extract_text() or ""
            return text.strip()
            
        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            workbook = openpyxl.load_workbook(uploaded_file.file)
            text = ""
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    text += row_text + "\n"
            return text.strip()
            
        elif filename.endswith(".docx"):
            doc = Document(uploaded_file.file)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            return text.strip()
            
        else:
            content = uploaded_file.file.read()
            return content.decode("utf-8", errors="ignore").strip()
            
    except Exception as e:
        print(f"Error extracting text: {e}")
        return ""


def chunk_text(text: str, chunk_size: int = 3000, chunk_overlap: int = 200):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap
    )
    return splitter.split_text(text)


def generate_ai_questions(context: str, num_questions: int, difficulty: str, question_type: str = "mixed"):
    """Generate questions using AI"""
    if not llm:
        raise Exception("LLM not initialized")
    
    if question_type == "mcq":
        types_instruction = "Generate ONLY multiple-choice questions with 4 options each."
    elif question_type == "short_answer":
        types_instruction = "Generate ONLY short answer questions."
    else:
        types_instruction = "Generate a MIX of multiple-choice (70%) and short answer (30%) questions."
    
    prompt = PromptTemplate(
        input_variables=["context", "num_questions", "difficulty", "types_instruction"],
        template="""You are an expert educator creating quiz questions.

Study Material:
{context}

Create {num_questions} questions at {difficulty} difficulty.
{types_instruction}

Return ONLY valid JSON:

[
  {{
    "question_text": "Question based on material...",
    "question_type": "multiple_choice",
    "code_snippet": null,
    "difficulty": "{difficulty}",
    "options": [
      {{"letter": "A", "text": "Option A"}},
      {{"letter": "B", "text": "Option B"}},
      {{"letter": "C", "text": "Option C"}},
      {{"letter": "D", "text": "Option D"}}
    ],
    "correct_answer": "A"
  }}
]
"""
    )
    
    response = llm.invoke(
        prompt.format(
            context=context[:3000],
            num_questions=num_questions,
            difficulty=difficulty.capitalize(),
            types_instruction=types_instruction
        )
    )
    
    content = response.content.strip()
    
    if content.startswith("```"):
        lines = content.split("\n")
        content = "\n".join(lines[1:-1])
        if content.startswith("json"):
            content = content[4:].strip()
    
    try:
        questions = json.loads(content)
        return questions
    except:
        start = content.find("[")
        end = content.rfind("]") + 1
        if start != -1 and end > start:
            try:
                questions = json.loads(content[start:end])
                return questions
            except:
                pass
        raise HTTPException(500, "AI generated invalid JSON")


def generate_mock_questions(num_questions: int, difficulty: str, content: str, question_type: str = "mixed"):
    """Generate mock questions"""
    words = content.split()[:100] if content else ["example"]
    questions = []
    
    for i in range(min(num_questions, 25)):
        if question_type == "short_answer" or (question_type == "mixed" and i % 3 == 2):
            questions.append({
                "question_text": f"Explain the concept of {words[i % len(words)]} mentioned in the material.",
                "question_type": "short_answer",
                "code_snippet": None,
                "difficulty": difficulty.capitalize(),
                "options": [],
                "correct_answer": f"Expected answer should explain {words[i % len(words)]} in detail."
            })
        else:
            questions.append({
                "question_text": f"Question {i+1}: Based on the material, which statement is correct?",
                "question_type": "multiple_choice",
                "code_snippet": None,
                "difficulty": difficulty.capitalize(),
                "options": [
                    {"letter": "A", "text": f"Statement about {words[(i*4) % len(words)]}"},
                    {"letter": "B", "text": f"Statement about {words[(i*4+1) % len(words)]}"},
                    {"letter": "C", "text": f"Statement about {words[(i*4+2) % len(words)]}"},
                    {"letter": "D", "text": f"Statement about {words[(i*4+3) % len(words)]}"}
                ],
                "correct_answer": "A"
            })
    
    return questions


def validate_questions(questions):
    valid_questions = []
    for q in questions:
        if "question_text" in q and "question_type" in q and "correct_answer" in q:
            if q["question_type"] == "multiple_choice":
                if "options" in q and len(q["options"]) == 4:
                    valid_questions.append(q)
            elif q["question_type"] == "short_answer":
                valid_questions.append(q)
    return valid_questions


def generate_pdf_review(results_data: dict) -> bytes:
    """Generate PDF with proper answer review"""
    buffer = io.BytesIO()
    
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72,
                           topMargin=72, bottomMargin=18)
    
    story = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#9333ea'),
        spaceAfter=30,
        alignment=TA_CENTER
    )
    
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=16,
        textColor=colors.HexColor('#9333ea'),
        spaceAfter=12,
        spaceBefore=12
    )
    
    story.append(Paragraph("Quiz Results Review", title_style))
    story.append(Spacer(1, 12))
    
    # Summary
    summary_data = [
        ['Score', f"{results_data['score_percentage']:.1f}%"],
        ['Correct Answers', f"{results_data['correct_answers']}/{results_data['total_questions']}"],
        ['Time Taken', f"{results_data['time_taken'] // 60}m {results_data['time_taken'] % 60}s"],
        ['Date', datetime.now().strftime('%Y-%m-%d %H:%M')]
    ]
    
    summary_table = Table(summary_data, colWidths=[2*inch, 3*inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f3f4f6')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 12),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.white)
    ]))
    
    story.append(summary_table)
    story.append(Spacer(1, 20))
    
    # Detailed Review
    story.append(Paragraph("Detailed Answer Review", heading_style))
    story.append(Spacer(1, 12))
    
    if 'detailed_results' in results_data and results_data['detailed_results']:
        for idx, item in enumerate(results_data['detailed_results'], 1):
            if item.get('is_correct') is True:
                status = "✓ Correct"
                status_color = colors.green
            elif item.get('is_correct') is False:
                status = "✗ Incorrect"
                status_color = colors.red
            else:
                status = "⚠ Needs Review"
                status_color = colors.orange
            
            q_type = item.get('question_type', 'multiple_choice')
            type_label = " (Short Answer)" if q_type == 'short_answer' else ""
            
            story.append(Paragraph(
                f"<b>Question {idx}{type_label}</b> - <font color='{status_color.hexval()}'>{status}</font>", 
                styles['Heading3']
            ))
            story.append(Spacer(1, 6))
            
            story.append(Paragraph(f"<b>Q:</b> {item['question_text']}", styles['Normal']))
            story.append(Spacer(1, 6))
            
            if q_type == 'multiple_choice':
                user_display = f"{item.get('user_answer', 'N/A')}. {item.get('user_answer_text', 'Not answered')}"
                correct_display = f"{item.get('correct_answer', '')}. {item.get('correct_answer_text', '')}"
            else:
                user_display = item.get('user_answer_text', 'Not answered')
                correct_display = item.get('correct_answer_text', 'Instructor will review')
            
            if item.get('is_correct') is True:
                story.append(Paragraph(
                    f"<b>Your Answer:</b> <font color='green'>{user_display}</font>", 
                    styles['Normal']
                ))
            else:
                story.append(Paragraph(
                    f"<b>Your Answer:</b> <font color='red'>{user_display}</font>", 
                    styles['Normal']
                ))
                story.append(Spacer(1, 4))
                
                if item.get('is_correct') is False:
                    story.append(Paragraph(
                        f"<b>Correct Answer:</b> <font color='green'>{correct_display}</font>", 
                        styles['Normal']
                    ))
                elif item.get('is_correct') is None:
                    story.append(Paragraph(
                        f"<b>Expected Answer:</b> {correct_display}", 
                        styles['Normal']
                    ))
            
            story.append(Spacer(1, 16))
    else:
        story.append(Paragraph("No detailed results available.", styles['Normal']))
    
    doc.build(story)
    
    pdf_bytes = buffer.getvalue()
    buffer.close()
    
    return pdf_bytes


# ==================== API ENDPOINTS ====================

@app.get("/")
async def root():
    return {
        "message": "NeuraNote Quiz API - With Levels",
        "status": "running",
        "ai_enabled": llm is not None,
        "features": ["quiz_levels", "difficulty_progression"]
    }


@app.get("/health")
async def health():
    return {
        "status": "healthy",
        "ai_configured": llm is not None
    }


@app.post("/api/quizzes/generate")
async def generate_quiz(
    files: List[UploadFile] = File(None),
    num_questions: int = Form(10),
    difficulty: str = Form("medium"),
    time_limit: int = Form(30),
    question_type: str = Form("mixed"),
    user_id: int = Form(1),
    note_id: Optional[int] = Form(None),
    source_content: Optional[str] = Form(None),  # NEW: For level progression
    db: Session = Depends(get_db)
):
    """Generate quiz - Enhanced with source content for levels"""
    
    if num_questions > 25:
        raise HTTPException(400, "Number of questions must not exceed 25")
    if num_questions < 1:
        raise HTTPException(400, "Number of questions must be at least 1")
    if time_limit < 1:
        raise HTTPException(400, "Time limit must be at least 1 minute")
    if time_limit > 180:
        raise HTTPException(400, "Time limit must not exceed 180 minutes")
    
    # Validate files only if no source_content provided
    if not source_content and files:
        for file in files:
            file.file.seek(0, 2)
            file_size = file.file.tell()
            file.file.seek(0)
            
            if file_size > 25 * 1024 * 1024:
                raise HTTPException(400, f"File {file.filename} exceeds 25MB limit")
    
    print(f"📝 Generating quiz: {num_questions} questions, {difficulty}, type: {question_type}")
    
    try:
        # Extract text from files OR use source_content
        if source_content:
            print("🔄 Using source content from previous level")
            combined_text = source_content
        else:
            if not files:
                raise HTTPException(400, "No files or source content provided")
            file_texts = []
            for file in files:
                text = extract_text_from_file(file)
                if text:
                    file_texts.append(text)
            
            if not file_texts:
                raise HTTPException(400, "No text could be extracted")
            
            combined_text = " ".join(file_texts)
        
        # Generate questions
        try:
            if llm:
                chunks = chunk_text(combined_text)
                context = " ".join(chunks[:2])
                questions_data = generate_ai_questions(context, num_questions, difficulty, question_type)
            else:
                raise Exception("No LLM")
        except Exception as e:
            print(f"⚠️ AI failed: {e}, using mock")
            questions_data = generate_mock_questions(num_questions, difficulty, combined_text, question_type)
        
        questions_data = validate_questions(questions_data)
        if not questions_data:
            raise HTTPException(500, "No valid questions generated")
        
        # Create quiz with source content stored
        quiz = models.Quiz(
            user_id=user_id,
            note_id=note_id,
            title=f"AI Quiz - {difficulty.capitalize()} - {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            description=f"Generated from {len(files) if not source_content else 'previous level'} file(s)",
            total_questions=len(questions_data),
            time_limit=time_limit * 60,
            difficulty=difficulty,
            source_content=combined_text  # Store for next level
        )
        db.add(quiz)
        db.flush()
        
        # Create questions
        for idx, q_data in enumerate(questions_data):
            question = models.Question(
                quiz_id=quiz.quiz_id,
                question_number=idx + 1,
                question_text=q_data['question_text'],
                code_snippet=q_data.get('code_snippet'),
                difficulty=q_data.get('difficulty', difficulty.capitalize()),
                question_type=q_data.get('question_type', 'multiple_choice')
            )
            db.add(question)
            db.flush()
            
            if q_data.get('question_type') == 'multiple_choice' and q_data.get('options'):
                for option_data in q_data['options']:
                    option = models.AnswerOption(
                        question_id=question.question_id,
                        option_letter=option_data['letter'],
                        option_text=option_data['text'],
                        is_correct=(option_data['letter'] == q_data['correct_answer'])
                    )
                    db.add(option)
            elif q_data.get('question_type') == 'short_answer':
                option = models.AnswerOption(
                    question_id=question.question_id,
                    option_letter='SA_EXPECTED',
                    option_text=q_data['correct_answer'],
                    is_correct=True
                )
                db.add(option)
        
        db.commit()
        db.refresh(quiz)
        
        questions = db.query(models.Question).filter(
            models.Question.quiz_id == quiz.quiz_id
        ).all()
        
        print(f"✅ Quiz created (ID: {quiz.quiz_id})")
        
        return {
            "quiz_id": quiz.quiz_id,
            "title": quiz.title,
            "description": quiz.description,
            "total_questions": quiz.total_questions,
            "time_limit": quiz.time_limit,
            "difficulty": quiz.difficulty,
            "source_content": combined_text,  # Return for frontend storage
            "questions": [
                {
                    "question_id": q.question_id,
                    "question_number": q.question_number,
                    "question_text": q.question_text,
                    "code_snippet": q.code_snippet,
                    "difficulty": q.difficulty,
                    "question_type": q.question_type,
                    "options": [
                        {
                            "option_id": opt.option_id,
                            "option_letter": opt.option_letter,
                            "option_text": opt.option_text
                        }
                        for opt in q.options if not opt.option_letter.startswith('SA_')
                    ] if q.question_type == 'multiple_choice' else []
                }
                for q in questions
            ]
        }
    
    except HTTPException:
        raise
    except Exception as e:
        print(f"❌ Error: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(500, str(e))


@app.post("/api/quizzes/{quiz_id}/submit")
async def submit_quiz(
    quiz_id: int,
    user_id: int = Form(...),
    answers: str = Form(...),
    time_taken: int = Form(...),
    db: Session = Depends(get_db)
):
    """Submit quiz and store detailed results"""
    
    print(f"📤 Submitting quiz {quiz_id}")
    print(f"📝 Answers received: {answers[:100]}...")
    
    try:
        answers_dict = json.loads(answers)
        print(f"📊 Parsed {len(answers_dict)} answers")
        
        questions = db.query(models.Question).filter(
            models.Question.quiz_id == quiz_id
        ).all()
        
        print(f"❓ Found {len(questions)} questions")
        
        if not questions:
            raise HTTPException(404, "Quiz not found")
        
        correct_answers = 0
        detailed_results = []
        
        for idx, question in enumerate(questions):
            user_answer = answers_dict.get(str(idx))
            print(f"Q{idx+1} ({question.question_type}): user_answer = {user_answer}")
            
            if question.question_type == 'multiple_choice':
                correct_option = db.query(models.AnswerOption).filter(
                    models.AnswerOption.question_id == question.question_id,
                    models.AnswerOption.is_correct == True
                ).first()
                
                is_correct = (user_answer == correct_option.option_letter) if user_answer else False
                if is_correct:
                    correct_answers += 1
                
                user_option = db.query(models.AnswerOption).filter(
                    models.AnswerOption.question_id == question.question_id,
                    models.AnswerOption.option_letter == user_answer
                ).first() if user_answer else None
                
                detailed_results.append({
                    "question_text": question.question_text,
                    "question_type": "multiple_choice",
                    "user_answer": user_answer or "N/A",
                    "user_answer_text": user_option.option_text if user_option else "Not answered",
                    "correct_answer": correct_option.option_letter,
                    "correct_answer_text": correct_option.option_text,
                    "is_correct": is_correct
                })
                
            elif question.question_type == 'short_answer':
                expected_option = db.query(models.AnswerOption).filter(
                    models.AnswerOption.question_id == question.question_id,
                    models.AnswerOption.option_letter.like('SA_%')
                ).first()
                
                print(f"✅ Short answer processed for Q{idx+1}")
                
                detailed_results.append({
                    "question_text": question.question_text,
                    "question_type": "short_answer",
                    "user_answer": user_answer or "",
                    "user_answer_text": user_answer or "Not answered",
                    "correct_answer": "Expected",
                    "correct_answer_text": expected_option.option_text if expected_option else "Instructor will review",
                    "is_correct": None
                })
        
        score_percentage = (correct_answers / len(questions)) * 100 if questions else 0
        print(f"💯 Score: {score_percentage}%, Correct: {correct_answers}/{len(questions)}")
        
        # Get quiz to check for next level
        quiz = db.query(models.Quiz).filter(models.Quiz.quiz_id == quiz_id).first()
        next_difficulty = get_next_difficulty(quiz.difficulty)
        
        # Save attempt with detailed results
        attempt = models.QuizAttempt(
            quiz_id=quiz_id,
            user_id=user_id,
            score_percentage=score_percentage,
            correct_answers=correct_answers,
            total_questions=len(questions),
            time_taken=time_taken,
            answers_json=json.dumps(detailed_results)
        )
        db.add(attempt)
        db.commit()
        db.refresh(attempt)
        
        print(f"✅ Quiz submitted. Score: {score_percentage}%")
        
        return {
            "attempt_id": attempt.attempt_id,
            "quiz_id": quiz_id,
            "score_percentage": score_percentage,
            "correct_answers": correct_answers,
            "total_questions": len(questions),
            "time_taken": time_taken,
            "detailed_results": detailed_results,
            "current_difficulty": quiz.difficulty,
            "next_difficulty": next_difficulty,
            "can_progress": next_difficulty is not None,
            "source_content": quiz.source_content if hasattr(quiz, 'source_content') else None
        }
    
    except json.JSONDecodeError as e:
        print(f"❌ JSON decode error: {e}")
        raise HTTPException(400, "Invalid answers format")
    except Exception as e:
        print(f"❌ Submission error: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(500, str(e))


@app.get("/api/quizzes/{quiz_id}/results/{attempt_id}/pdf")
async def download_results_pdf(
    quiz_id: int, 
    attempt_id: int,
    db: Session = Depends(get_db)
):
    """Download PDF"""
    
    try:
        attempt = db.query(models.QuizAttempt).filter(
            models.QuizAttempt.attempt_id == attempt_id,
            models.QuizAttempt.quiz_id == quiz_id
        ).first()
        
        if not attempt:
            raise HTTPException(404, "Quiz attempt not found")
        
        detailed_results = []
        if hasattr(attempt, 'answers_json') and attempt.answers_json:
            try:
                detailed_results = json.loads(attempt.answers_json)
            except:
                pass
        
        results_data = {
            "score_percentage": attempt.score_percentage,
            "correct_answers": attempt.correct_answers,
            "total_questions": attempt.total_questions,
            "time_taken": attempt.time_taken,
            "detailed_results": detailed_results
        }
        
        pdf_bytes = generate_pdf_review(results_data)
        
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename=quiz_results_{quiz_id}_{attempt_id}.pdf"
            }
        )
        
    except Exception as e:
        print(f"❌ PDF Error: {e}")
        import traceback
        traceback.print_exc()
        raise HTTPException(500, str(e))


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)

from typing import Dict, Any
from sqlalchemy import func, desc

# QUIZ HISTORY & ANALYTICS ENDPOINTS

@app.get("/api/quizzes/history/{user_id}")
async def get_quiz_history(
    user_id: int,
    limit: int = 10,
    offset: int = 0,
    db: Session = Depends(get_db)
):
    """
    Get quiz attempt history for a user with pagination
    
    Returns:
    - List of quiz attempts with quiz details
    - Total count of attempts
    - Pagination info
    """
    
    # Get total count
    total_count = db.query(models.QuizAttempt).filter(
        models.QuizAttempt.user_id == user_id
    ).count()
    
    # Get attempts with quiz details
    attempts = db.query(
        models.QuizAttempt,
        models.Quiz
    ).join(
        models.Quiz, models.QuizAttempt.quiz_id == models.Quiz.quiz_id
    ).filter(
        models.QuizAttempt.user_id == user_id
    ).order_by(
        desc(models.QuizAttempt.attempt_date)
    ).limit(limit).offset(offset).all()
    
    history = []
    for attempt, quiz in attempts:
        history.append({
            "attempt_id": attempt.attempt_id,
            "quiz_id": quiz.quiz_id,
            "quiz_title": quiz.title,
            "quiz_description": quiz.description,
            "difficulty": quiz.difficulty,
            "total_questions": attempt.total_questions,
            "correct_answers": attempt.correct_answers,
            "score_percentage": float(attempt.score_percentage) if attempt.score_percentage else 0,
            "time_taken": attempt.time_taken,
            "time_limit": quiz.time_limit,
            "attempt_date": attempt.attempt_date.isoformat() if attempt.attempt_date else None,
            "passed": float(attempt.score_percentage) >= 70 if attempt.score_percentage else False
        })
    
    return {
        "history": history,
        "total_count": total_count,
        "limit": limit,
        "offset": offset,
        "has_more": (offset + limit) < total_count
    }


@app.get("/api/quizzes/analytics/{user_id}")
async def get_quiz_analytics(
    user_id: int,
    db: Session = Depends(get_db)
):
    """
    Get comprehensive analytics for a user's quiz performance
    
    Returns:
    - Overall statistics
    - Performance by difficulty
    - Recent trends
    - Best/worst performances
    """
    
    # Total attempts
    total_attempts = db.query(func.count(models.QuizAttempt.attempt_id)).filter(
        models.QuizAttempt.user_id == user_id
    ).scalar() or 0
    
    if total_attempts == 0:
        return {
            "total_attempts": 0,
            "overall_stats": {},
            "by_difficulty": {},
            "recent_trend": [],
            "best_attempts": [],
            "time_stats": {}
        }
    
    # Overall statistics
    overall_stats = db.query(
        func.avg(models.QuizAttempt.score_percentage).label('avg_score'),
        func.max(models.QuizAttempt.score_percentage).label('best_score'),
        func.min(models.QuizAttempt.score_percentage).label('worst_score'),
        func.sum(models.QuizAttempt.correct_answers).label('total_correct'),
        func.sum(models.QuizAttempt.total_questions).label('total_questions')
    ).filter(
        models.QuizAttempt.user_id == user_id
    ).first()
    
    # Performance by difficulty
    difficulty_stats = db.query(
        models.Quiz.difficulty,
        func.count(models.QuizAttempt.attempt_id).label('attempts'),
        func.avg(models.QuizAttempt.score_percentage).label('avg_score'),
        func.max(models.QuizAttempt.score_percentage).label('best_score')
    ).join(
        models.Quiz, models.QuizAttempt.quiz_id == models.Quiz.quiz_id
    ).filter(
        models.QuizAttempt.user_id == user_id
    ).group_by(
        models.Quiz.difficulty
    ).all()
    
    by_difficulty = {}
    for difficulty, attempts, avg_score, best_score in difficulty_stats:
        by_difficulty[difficulty] = {
            "attempts": attempts,
            "average_score": float(avg_score) if avg_score else 0,
            "best_score": float(best_score) if best_score else 0
        }
    
    # Recent trend (last 7 attempts)
    recent_attempts = db.query(
        models.QuizAttempt.score_percentage,
        models.QuizAttempt.attempt_date,
        models.Quiz.difficulty
    ).join(
        models.Quiz, models.QuizAttempt.quiz_id == models.Quiz.quiz_id
    ).filter(
        models.QuizAttempt.user_id == user_id
    ).order_by(
        desc(models.QuizAttempt.attempt_date)
    ).limit(7).all()
    
    recent_trend = [
        {
            "score": float(score) if score else 0,
            "date": date.isoformat() if date else None,
            "difficulty": difficulty
        }
        for score, date, difficulty in reversed(recent_attempts)
    ]
    
    # Best attempts (top 5)
    best_attempts_query = db.query(
        models.QuizAttempt,
        models.Quiz
    ).join(
        models.Quiz, models.QuizAttempt.quiz_id == models.Quiz.quiz_id
    ).filter(
        models.QuizAttempt.user_id == user_id
    ).order_by(
        desc(models.QuizAttempt.score_percentage),
        models.QuizAttempt.time_taken
    ).limit(5).all()
    
    best_attempts = [
        {
            "quiz_title": quiz.title,
            "difficulty": quiz.difficulty,
            "score": float(attempt.score_percentage) if attempt.score_percentage else 0,
            "time_taken": attempt.time_taken,
            "date": attempt.attempt_date.isoformat() if attempt.attempt_date else None
        }
        for attempt, quiz in best_attempts_query
    ]
    
    # Time statistics
    time_stats = db.query(
        func.avg(models.QuizAttempt.time_taken).label('avg_time'),
        func.min(models.QuizAttempt.time_taken).label('fastest_time'),
        func.max(models.QuizAttempt.time_taken).label('slowest_time')
    ).filter(
        models.QuizAttempt.user_id == user_id
    ).first()
    
    # Pass rate
    passed_count = db.query(func.count(models.QuizAttempt.attempt_id)).filter(
        models.QuizAttempt.user_id == user_id,
        models.QuizAttempt.score_percentage >= 70
    ).scalar() or 0
    
    return {
        "total_attempts": total_attempts,
        "overall_stats": {
            "average_score": float(overall_stats.avg_score) if overall_stats.avg_score else 0,
            "best_score": float(overall_stats.best_score) if overall_stats.best_score else 0,
            "worst_score": float(overall_stats.worst_score) if overall_stats.worst_score else 0,
            "total_correct": overall_stats.total_correct or 0,
            "total_questions": overall_stats.total_questions or 0,
            "pass_rate": (passed_count / total_attempts * 100) if total_attempts > 0 else 0
        },
        "by_difficulty": by_difficulty,
        "recent_trend": recent_trend,
        "best_attempts": best_attempts,
        "time_stats": {
            "average_time": time_stats.avg_time or 0,
            "fastest_time": time_stats.fastest_time or 0,
            "slowest_time": time_stats.slowest_time or 0
        }
    }


@app.get("/api/quizzes/attempt/{attempt_id}/details")
async def get_attempt_details(
    attempt_id: int,
    user_id: int,
    db: Session = Depends(get_db)
):
    """
    Get detailed information about a specific quiz attempt
    Including all questions and answers
    """
    
    # Get attempt with quiz
    attempt = db.query(models.QuizAttempt).filter(
        models.QuizAttempt.attempt_id == attempt_id,
        models.QuizAttempt.user_id == user_id
    ).first()
    
    if not attempt:
        raise HTTPException(status_code=404, detail="Attempt not found")
    
    quiz = db.query(models.Quiz).filter(
        models.Quiz.quiz_id == attempt.quiz_id
    ).first()
    
    # Parse answers from JSON
    detailed_results = []
    if attempt.answers_json:
        try:
            detailed_results = json.loads(attempt.answers_json)
        except:
            # Fallback if JSON parsing fails
            detailed_results = []
    
    return {
        "attempt_id": attempt.attempt_id,
        "quiz_id": quiz.quiz_id,
        "quiz_title": quiz.title,
        "quiz_description": quiz.description,
        "difficulty": quiz.difficulty,
        "total_questions": attempt.total_questions,
        "correct_answers": attempt.correct_answers,
        "score_percentage": float(attempt.score_percentage) if attempt.score_percentage else 0,
        "time_taken": attempt.time_taken,
        "time_limit": quiz.time_limit,
        "attempt_date": attempt.attempt_date.isoformat() if attempt.attempt_date else None,
        "detailed_results": detailed_results
    }


@app.delete("/api/quizzes/history/{attempt_id}")
async def delete_quiz_attempt(
    attempt_id: int,
    user_id: int,
    db: Session = Depends(get_db)
):
    """
    Delete a specific quiz attempt from history
    """
    
    attempt = db.query(models.QuizAttempt).filter(
        models.QuizAttempt.attempt_id == attempt_id,
        models.QuizAttempt.user_id == user_id
    ).first()
    
    if not attempt:
        raise HTTPException(status_code=404, detail="Attempt not found")
    
    db.delete(attempt)
    db.commit()
    
    return {"message": "Attempt deleted successfully", "attempt_id": attempt_id}
